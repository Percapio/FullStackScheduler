"""Shipping log generation (Phase 33 §6.1-6.3)."""
from __future__ import annotations

import io
import logging
import sys
import time
import zipfile
from copy import copy
from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.orm import Session

from backend.app.api.deps import get_wall_clock
from backend.app.config import Settings, bundled_resource_path, get_settings
from backend.app.models import Assembly, Base, BuildType, Customer, Job, JobStatus
from backend.app.services import shipping_log as shipping_log_service
from backend.app.services.history_export import flatten_operator_notes, strip_operator_markup
from backend.app.services.shipping_log import (
    EXCEL_CELL_MAX_CHARS,
    SHIPPING_LOG_TEMPLATE_FILENAME,
    SHIPPING_LOG_TEMPLATE_LAYOUT,
    TRUNCATION_MARKER,
    IneligibleReason,
    TemplateReady,
    TemplateUnavailable,
    TemplateUnavailableKind,
    WorkbookStampingError,
    load_shipping_log_template,
    plan_shipping_log_layout,
    render_box,
    resolve_log_jobs,
    stamp_workbook,
    text_cell_value,
)

TEMPLATE_PATH = bundled_resource_path(SHIPPING_LOG_TEMPLATE_FILENAME)
SHEET = SHIPPING_LOG_TEMPLATE_LAYOUT.sheet_title
PINNED_UTC = datetime(2026, 9, 16, 12, 30, 45, tzinfo=timezone.utc)


# ---- helpers -----------------------------------------------------------------


def _add_job(
    session: Session,
    *,
    part_number: str,
    notes: str | None = None,
    customer_name: str = "ACME",
    **overrides,
) -> Job:
    assembly = Assembly(part_number=part_number, base_mfg_notes=notes)
    session.add(assembly)
    session.flush()
    customer = session.scalars(select(Customer).where(Customer.name == customer_name)).first()
    if customer is None:
        customer = Customer(name=customer_name)
        session.add(customer)
        session.flush()
    job_fields = dict(
        assembly_id=assembly.id,
        customer_id=customer.id,
        quantity=10,
        build_type=BuildType.new,
        status=JobStatus.planned,
    )
    job_fields.update(overrides)
    job = Job(**job_fields)
    session.add(job)
    session.flush()
    return job


def _bundled_template() -> TemplateReady:
    template = load_shipping_log_template(TEMPLATE_PATH, SHIPPING_LOG_TEMPLATE_LAYOUT)
    assert isinstance(template, TemplateReady), template
    return template


def _template_sheet():
    return load_workbook(TEMPLATE_PATH)[SHEET]


def _sheet_of(content: bytes):
    return load_workbook(io.BytesIO(content))[SHEET]


def _resolved_style(cell) -> tuple:
    # copy() unwraps openpyxl's StyleProxy; two proxies never compare equal.
    return (
        copy(cell.font), copy(cell.border), copy(cell.fill),
        cell.number_format, copy(cell.alignment), copy(cell.protection),
    )


def _use_settings(client, **overrides) -> Settings:
    settings = Settings(**overrides)
    client.app.dependency_overrides[get_settings] = lambda: settings
    return settings


def _generate(client, job_ids: list[int]):
    return client.post("/api/shipping-log", json={"job_ids": job_ids})


def _box_start(box_index: int, gap_rows: int = 1) -> int:
    return 1 + box_index * (SHIPPING_LOG_TEMPLATE_LAYOUT.box_rows + gap_rows)


def _write_modified_template(tmp_path: Path, edit) -> Path:
    workbook = load_workbook(TEMPLATE_PATH)
    edit(workbook)
    path = tmp_path / "modified_template.xlsx"
    workbook.save(path)
    return path


@pytest.fixture()
def shipping_log_logger(monkeypatch):
    # Another test's logging.config.fileConfig can disable existing loggers.
    api_logger = logging.getLogger("backend.app.api.shipping_log")
    service_logger = logging.getLogger("backend.app.services.shipping_log")
    monkeypatch.setattr(api_logger, "disabled", False)
    monkeypatch.setattr(service_logger, "disabled", False)
    return api_logger


# ---- §6.1 template -----------------------------------------------------------


def test_bundled_template_loads_ready_with_measured_heights():
    template = _bundled_template()
    sheet = _template_sheet()
    default = sheet.sheet_format.defaultRowHeight
    declared = [
        sheet.row_dimensions[row].height or default
        for row in range(1, SHIPPING_LOG_TEMPLATE_LAYOUT.box_rows + 1)
    ]

    assert template.box_height_points == pytest.approx(sum(declared))
    assert template.box_height_points == pytest.approx(111.0)
    assert template.notes_row_height_points == pytest.approx(35.25)
    assert template.notes_available_points == pytest.approx(35.25 + 15 + 15 + 15)
    assert template.default_row_height_points == pytest.approx(15.0)
    assert len(template.sha256) == 64


@pytest.mark.parametrize(
    ("edit", "expected_kind"),
    [
        (lambda wb: setattr(wb[SHEET], "title", "OTHER"), TemplateUnavailableKind.SHEET_MISSING),
        (lambda wb: wb[SHEET].__setitem__("A1", "PART"), TemplateUnavailableKind.HEADER_MISMATCH),
        (lambda wb: wb[SHEET].unmerge_cells("E2:E5"), TemplateUnavailableKind.NOTES_MERGE_MISSING),
        (lambda wb: wb[SHEET].__setitem__("A8", "extra"), TemplateUnavailableKind.EXTRA_CONTENT_BELOW_BOX),
        (lambda wb: wb[SHEET].merge_cells("G8:G9"), TemplateUnavailableKind.EXTRA_CONTENT_BELOW_BOX),
    ],
    ids=["renamed-sheet", "a1-changed", "notes-unmerged", "content-in-a8", "merge-below-box"],
)
def test_structural_defects_are_unavailable(tmp_path, edit, expected_kind):
    path = _write_modified_template(tmp_path, edit)

    template = load_shipping_log_template(path, SHIPPING_LOG_TEMPLATE_LAYOUT)

    assert isinstance(template, TemplateUnavailable)
    assert template.kind is expected_kind
    assert template.sha256 is not None


def test_header_mismatch_names_the_cell_and_found_text(tmp_path):
    path = _write_modified_template(tmp_path, lambda wb: wb[SHEET].__setitem__("A1", "PART"))

    template = load_shipping_log_template(path, SHIPPING_LOG_TEMPLATE_LAYOUT)

    assert (template.cell, template.found) == ("A1", "PART")


def test_missing_file_is_unavailable(tmp_path):
    template = load_shipping_log_template(tmp_path / "absent.xlsx", SHIPPING_LOG_TEMPLATE_LAYOUT)
    assert template == TemplateUnavailable(TemplateUnavailableKind.MISSING)


def test_unreadable_path_is_unavailable(tmp_path):
    template = load_shipping_log_template(tmp_path, SHIPPING_LOG_TEMPLATE_LAYOUT)
    assert template == TemplateUnavailable(TemplateUnavailableKind.UNREADABLE)


def test_non_zip_bytes_are_not_a_workbook(tmp_path):
    path = tmp_path / "junk.xlsx"
    path.write_bytes(b"this is not a zip archive")

    template = load_shipping_log_template(path, SHIPPING_LOG_TEMPLATE_LAYOUT)

    assert isinstance(template, TemplateUnavailable)
    assert template.kind is TemplateUnavailableKind.NOT_A_WORKBOOK
    assert template.sha256 is not None


def test_frozen_build_resolves_resources_under_meipass(monkeypatch, tmp_path):
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(sys, "_MEIPASS", str(tmp_path), raising=False)

    resolved = bundled_resource_path(SHIPPING_LOG_TEMPLATE_FILENAME)

    assert resolved == tmp_path / "backend" / "app" / "resources" / SHIPPING_LOG_TEMPLATE_FILENAME


def test_dev_build_resolves_resources_in_the_package():
    assert TEMPLATE_PATH.is_file()
    assert TEMPLATE_PATH.parent.name == "resources"


def test_app_starting_without_template_reports_it_and_answers_503(
    monkeypatch, tmp_path, session_factory, caplog, shipping_log_logger,
):
    from fastapi.testclient import TestClient
    from backend.app.api import create_app
    from backend.app.api.deps import get_session

    missing_path = tmp_path / "nowhere" / SHIPPING_LOG_TEMPLATE_FILENAME
    monkeypatch.setattr(shipping_log_service, "bundled_resource_path", lambda name: missing_path)
    app = create_app()

    def _session():
        session = session_factory()
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_session] = _session
    caplog.set_level(logging.ERROR, logger="backend.app.services.shipping_log")
    with TestClient(app) as client:
        candidates = client.get("/api/shipping-log/candidates")
        generated = _generate(client, [1])

    assert candidates.status_code == 200
    assert candidates.json()["template_ready"] is False
    assert generated.status_code == 503
    assert generated.json() == {"kind": "template_unavailable"}
    assert "nowhere" not in generated.text
    startup_errors = [r for r in caplog.records if r.name == "backend.app.services.shipping_log"]
    assert len(startup_errors) == 1
    assert "missing" in startup_errors[0].getMessage()


# ---- §6.2 candidates ---------------------------------------------------------


def _seed_mixed_population(session: Session) -> dict[str, Job]:
    jobs = {
        "late": _add_job(session, part_number="POP-LATE", resolved_ship_date=date(2026, 10, 1)),
        "undated": _add_job(session, part_number="POP-UNDATED", resolved_ship_date=None, ship_date_text="???"),
        "early": _add_job(session, part_number="POP-EARLY", resolved_ship_date=date(2026, 9, 1)),
        "shipped": _add_job(session, part_number="POP-SHIPPED", status=JobStatus.shipped,
                            resolved_ship_date=date(2026, 8, 1), shipped_at=date(2026, 8, 1)),
        "discarded": _add_job(session, part_number="POP-DISCARDED", resolved_ship_date=date(2026, 9, 2),
                              discarded_at=datetime(2026, 9, 3)),
        "superseded": _add_job(session, part_number="POP-SUPERSEDED", resolved_ship_date=date(2026, 9, 2),
                               superseded_at=datetime(2026, 9, 3)),
        "undated_second": _add_job(session, part_number="POP-UNDATED-2", resolved_ship_date=None),
    }
    session.commit()
    return jobs


def test_candidates_equal_the_shipping_view_population_and_order(client, session):
    _seed_mixed_population(session)

    shipping_ids = [job["id"] for job in client.get("/api/jobs/shipping").json()]
    candidate_ids = [c["job_id"] for c in client.get("/api/shipping-log/candidates").json()["candidates"]]

    assert candidate_ids == shipping_ids


def test_candidates_exclude_superseded_discarded_and_shipped(client, session):
    jobs = _seed_mixed_population(session)

    body = client.get("/api/shipping-log/candidates").json()
    candidate_ids = {c["job_id"] for c in body["candidates"]}

    assert candidate_ids == {jobs[k].id for k in ("late", "undated", "early", "undated_second")}
    assert body["total"] == 4
    assert body["truncated"] is False
    assert body["template_ready"] is True
    assert body["max_jobs_per_log"] == Settings().shipping_log_max_jobs


def test_null_ship_date_sorts_last(client, session):
    jobs = _seed_mixed_population(session)

    candidates = client.get("/api/shipping-log/candidates").json()["candidates"]

    assert [c["job_id"] for c in candidates] == [
        jobs["early"].id, jobs["late"].id, jobs["undated"].id, jobs["undated_second"].id,
    ]
    assert candidates[2]["ship_date_text"] == "???"


def test_candidate_carries_the_full_identity(client, session):
    from backend.app.models import BuildQualifier

    _add_job(
        session, part_number="138537", customer_name="Widget Co", split_suffix="-bal",
        build_type=BuildType.ronc, repeat_reference="137001", build_qualifier=BuildQualifier.rwk,
        quantity=40, resolved_ship_date=date(2026, 9, 30),
    )
    session.commit()

    candidate = client.get("/api/shipping-log/candidates").json()["candidates"][0]

    assert candidate == {
        "job_id": candidate["job_id"], "part_number": "138537", "split_suffix": "-bal",
        "build_type": "ronc", "repeat_reference": "137001", "build_qualifier": "rwk",
        "quantity": 40, "resolved_ship_date": "2026-09-30", "ship_date_text": None,
        "customer_name": "Widget Co",
    }


def test_population_over_the_cap_is_truncated_not_hidden(client, session):
    for index in range(5):
        _add_job(session, part_number=f"CAP-{index}", resolved_ship_date=date(2026, 9, 1 + index))
    session.commit()
    _use_settings(client, shipping_log_candidate_max=3)

    body = client.get("/api/shipping-log/candidates").json()

    assert len(body["candidates"]) == 3
    assert body["total"] == 5
    assert body["truncated"] is True


# ---- §6.3 generation: layout and styling -------------------------------------


def test_two_jobs_give_two_styled_boxes_with_their_own_merges(client, session):
    first = _add_job(session, part_number="012345", resolved_ship_date=date(2026, 9, 1))
    second = _add_job(session, part_number="067890", resolved_ship_date=date(2026, 9, 2))
    session.commit()

    response = _generate(client, [first.id, second.id])

    assert response.status_code == 200
    sheet = _sheet_of(response.content)
    template = _template_sheet()
    assert {merged.coord for merged in sheet.merged_cells.ranges} == {"E2:E5", "E9:E12"}
    assert sheet["A2"].value == "012345"
    assert sheet["A9"].value == "067890"
    assert (sheet["B2"].value, sheet["B9"].value) == (10, 10)
    for box_index in (0, 1):
        offset = _box_start(box_index) - 1
        for row in range(1, 7):
            for column in range(1, 6):
                stamped = sheet.cell(row=row + offset, column=column)
                original = template.cell(row=row, column=column)
                assert _resolved_style(stamped) == _resolved_style(original), stamped.coordinate
            if row != SHIPPING_LOG_TEMPLATE_LAYOUT.notes_row:
                assert sheet.row_dimensions[row + offset].height == template.row_dimensions[row].height


def test_column_widths_orientation_and_margins_are_the_templates(client, session):
    job = _add_job(session, part_number="PAGE-1")
    session.commit()

    sheet = _sheet_of(_generate(client, [job.id]).content)
    template = _template_sheet()

    for letter in "ABCDE":
        assert sheet.column_dimensions[letter].width == template.column_dimensions[letter].width
    assert sheet.page_setup.orientation == "landscape"
    assert (sheet.page_margins.left, sheet.page_margins.top, sheet.page_margins.bottom) == (0.25, 0.75, 0.75)


def test_every_box_keeps_the_template_labels_and_an_empty_ship_type_grid(client, session):
    jobs = [_add_job(session, part_number=f"GRID-{i}", notes=f"note {i}", resolved_ship_date=date(2026, 9, 1 + i))
            for i in range(3)]
    session.commit()

    sheet = _sheet_of(_generate(client, [job.id for job in jobs]).content)
    template = _template_sheet()

    for box_index in range(3):
        offset = _box_start(box_index) - 1
        for cell_ref, label in SHIPPING_LOG_TEMPLATE_LAYOUT.header_expectations:
            assert sheet[cell_ref].offset(row=offset).value == label
        for row in range(3, 7):
            for column in range(1, 5):
                assert sheet.cell(row=row + offset, column=column).value == template.cell(row=row, column=column).value
        assert sheet.cell(row=3 + offset, column=4).value is None
        assert sheet.cell(row=2 + offset, column=3).value is None


def test_copy_pass_runs_before_any_job_value_is_written(client, session):
    with_notes = _add_job(session, part_number="ORDER-0", notes="ship with care", resolved_ship_date=date(2026, 9, 1))
    without_notes = _add_job(session, part_number="ORDER-1", notes=None, resolved_ship_date=date(2026, 9, 2))
    session.commit()

    sheet = _sheet_of(_generate(client, [with_notes.id, without_notes.id]).content)

    assert sheet["E2"].value == "ship with care"
    assert sheet["E9"].value is None
    assert sheet["A9"].value == "ORDER-1"


def test_boxes_follow_shipping_order_not_submission_order(client, session):
    late = _add_job(session, part_number="SEQ-LATE", resolved_ship_date=date(2026, 9, 30))
    undated = _add_job(session, part_number="SEQ-UNDATED", resolved_ship_date=None)
    early = _add_job(session, part_number="SEQ-EARLY", resolved_ship_date=date(2026, 9, 1))
    session.commit()

    sheet = _sheet_of(_generate(client, [undated.id, late.id, early.id]).content)

    assert [sheet.cell(row=_box_start(i) + 1, column=1).value for i in range(3)] == [
        "SEQ-EARLY", "SEQ-LATE", "SEQ-UNDATED",
    ]


def test_duplicate_ids_give_one_box_each(client, session):
    job = _add_job(session, part_number="DUP-1")
    session.commit()

    sheet = _sheet_of(_generate(client, [job.id, job.id, job.id]).content)

    assert {merged.coord for merged in sheet.merged_cells.ranges} == {"E2:E5"}
    assert sheet["A9"].value is None


# ---- §6.3 generation: cell safety --------------------------------------------


def test_formula_notes_are_stored_as_neutralised_text(client, session):
    job = _add_job(session, part_number="=HYPERLINK(1)", notes="=cmd|' /C calc'!A0")
    session.commit()

    response = _generate(client, [job.id])

    sheet = _sheet_of(response.content)
    assert sheet["E2"].data_type == "s"
    assert sheet["E2"].value == "'=cmd|' /C calc'!A0"
    assert sheet["A2"].data_type == "s"
    assert sheet["A2"].value.startswith("'")
    sheet_xml = zipfile.ZipFile(io.BytesIO(response.content)).read("xl/worksheets/sheet1.xml").decode()
    assert "<f>" not in sheet_xml and "<f " not in sheet_xml


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@", "\t", "\r"])
def test_every_formula_prefix_is_neutralised(prefix):
    assert text_cell_value(f"{prefix}SUM(A1)").text == f"'{prefix}SUM(A1)"


@pytest.mark.parametrize("prefix", ["+", "-", "@"])
def test_rendered_notes_and_part_numbers_carry_the_neutralised_prefix(session, prefix):
    job = _add_job(session, part_number=f"{prefix}PN", notes=f"{prefix}SUM(A1)")

    box = render_box(job, Settings())

    assert box.notes.text == f"'{prefix}SUM(A1)"
    assert box.part_number.text == f"'{prefix}PN"


def test_note_lines_are_trimmed_so_whitespace_prefixes_never_lead(session):
    job = _add_job(session, part_number="TRIM-1", notes="\t=SUM(A1)")

    box = render_box(job, Settings())

    assert box.notes.text == "'=SUM(A1)"


def test_error_code_text_is_not_written_as_an_error_cell(client, session):
    job = _add_job(session, part_number="ERR-1", notes="#N/A")
    session.commit()

    sheet = _sheet_of(_generate(client, [job.id]).content)

    assert (sheet["E2"].data_type, sheet["E2"].value) == ("s", "#N/A")


def test_control_characters_are_removed_instead_of_failing_the_log(client, session):
    job = _add_job(session, part_number="CTRL-1", notes="bell\x07 here\x01")
    session.commit()

    response = _generate(client, [job.id])

    assert response.status_code == 200
    assert _sheet_of(response.content)["E2"].value == "bell here"


def test_notes_markup_is_stripped_and_line_breaks_kept(client, session):
    job = _add_job(session, part_number="MARK-1", notes="**First** line\n~~removed~~\n\n  *second*  \n\nthird")
    session.commit()

    sheet = _sheet_of(_generate(client, [job.id]).content)

    assert sheet["E2"].value == "First line\nsecond\nthird"


def test_part_number_leading_zeros_survive(client, session):
    job = _add_job(session, part_number="000123")
    session.commit()

    sheet = _sheet_of(_generate(client, [job.id]).content)

    assert sheet["A2"].value == "000123"
    assert sheet["A2"].number_format == "@"


def test_flatten_operator_notes_is_the_shared_tokenisation_joined_with_pipes():
    raw = "**a**\n~~gone~~\n\n b \n~~unpaired"
    assert strip_operator_markup(raw) == ["a", "b", "~~unpaired"]
    assert flatten_operator_notes(raw) == "a | b | ~~unpaired"
    assert strip_operator_markup(None) == []
    assert flatten_operator_notes(None) == ""


# ---- §6.3 generation: clipping -----------------------------------------------


def test_notes_over_the_cell_limit_are_truncated_and_reported(client, session):
    job = _add_job(session, part_number="LONG-1", notes="x" * 40_000)
    session.commit()

    response = _generate(client, [job.id])

    assert response.status_code == 200
    assert response.headers["x-shipping-log-clipped"] == str(job.id)
    notes = _sheet_of(response.content)["E2"].value
    assert notes.endswith(TRUNCATION_MARKER)
    assert len(notes) <= EXCEL_CELL_MAX_CHARS


def test_cell_limit_counts_utf16_units(session):
    job = _add_job(session, part_number="EMOJI-1", notes="\U0001F4E6" * 20_000)

    box = render_box(job, Settings())

    assert box.notes_clipped is True
    assert len(box.notes.text.encode("utf-16-le")) // 2 <= EXCEL_CELL_MAX_CHARS
    assert box.notes.text.endswith(TRUNCATION_MARKER)


def test_notes_taller_than_the_row_ceiling_are_capped_and_reported(client, session):
    tall = _add_job(session, part_number="TALL-1", notes="\n".join(f"line {i}" for i in range(40)),
                    resolved_ship_date=date(2026, 9, 1))
    short = _add_job(session, part_number="SHORT-1", notes="one line", resolved_ship_date=date(2026, 9, 2))
    session.commit()

    response = _generate(client, [tall.id, short.id])

    sheet = _sheet_of(response.content)
    assert sheet.row_dimensions[2].height == 409
    assert sheet.row_dimensions[9].height == 35.25
    assert response.headers["x-shipping-log-clipped"] == str(tall.id)


def test_notes_row_grows_with_the_line_estimate(session):
    job = _add_job(session, part_number="GROW-1", notes="\n".join(["short"] * 8))
    template = _bundled_template()

    plan = plan_shipping_log_layout(template, [render_box(job, Settings())], Settings())

    growth = 8 * 15 - template.notes_available_points
    assert plan.placements[0].notes_row_points == pytest.approx(35.25 + growth)
    assert plan.placements[0].notes_height_clipped is False


def test_no_clipped_header_when_nothing_is_clipped(client, session):
    job = _add_job(session, part_number="FITS-1", notes="fits")
    session.commit()

    assert "x-shipping-log-clipped" not in _generate(client, [job.id]).headers


# ---- §6.3 generation: page breaks --------------------------------------------


def _plain_boxes(session: Session, count: int) -> list:
    return [render_box(_add_job(session, part_number=f"PB-{count}-{i}"), Settings()) for i in range(count)]


def test_five_boxes_break_after_the_fourth_counting_gap_rows(session):
    template = _bundled_template()
    settings = Settings(shipping_log_box_gap_rows=1, shipping_log_page_body_points=504)

    plan = plan_shipping_log_layout(template, _plain_boxes(session, 5), settings)

    assert [p.height_points for p in plan.placements] == [111.0] * 5
    assert plan.row_breaks_after == [_box_start(3) + 5]
    assert plan.row_breaks_after == [27]


def test_without_gap_rows_the_fifth_box_still_breaks(session):
    template = _bundled_template()
    settings = Settings(shipping_log_box_gap_rows=0, shipping_log_page_body_points=504)

    plan = plan_shipping_log_layout(template, _plain_boxes(session, 5), settings)

    assert [p.start_row for p in plan.placements] == [1, 7, 13, 19, 25]
    assert plan.row_breaks_after == [24]


def test_box_taller_than_a_page_sits_alone_with_no_break_inside(session):
    template = _bundled_template()
    settings = Settings(shipping_log_page_body_points=200)
    boxes = [
        render_box(_add_job(session, part_number="HUGE-0"), settings),
        render_box(_add_job(session, part_number="HUGE-1", notes="\n".join(["x"] * 40)), settings),
        render_box(_add_job(session, part_number="HUGE-2"), settings),
    ]

    plan = plan_shipping_log_layout(template, boxes, settings)

    assert plan.placements[1].height_points > settings.shipping_log_page_body_points
    assert plan.row_breaks_after == [6, 13]
    assert not any(8 <= row < 13 for row in plan.row_breaks_after)


def test_page_breaks_reach_the_saved_workbook(client, session):
    jobs = [_add_job(session, part_number=f"BRK-{i}", resolved_ship_date=date(2026, 9, 1 + i)) for i in range(5)]
    session.commit()

    sheet = _sheet_of(_generate(client, [job.id for job in jobs]).content)

    assert [brk.id for brk in sheet.row_breaks.brk] == [27]


# ---- §6.3 generation: failures -----------------------------------------------


def test_empty_selection_is_a_selection_size_422(client):
    response = _generate(client, [])

    assert response.status_code == 422
    assert response.json() == {"kind": "selection_size", "requested": 0, "max": 200}


def test_selection_over_max_is_a_selection_size_422(client):
    max_jobs = Settings().shipping_log_max_jobs

    response = _generate(client, list(range(1, max_jobs + 2)))

    assert response.status_code == 422
    assert response.json() == {"kind": "selection_size", "requested": max_jobs + 1, "max": max_jobs}


def test_array_over_the_parse_ceiling_is_rejected_by_the_schema(client):
    max_jobs = Settings().shipping_log_max_jobs

    response = _generate(client, list(range(1, max_jobs * 4 + 2)))

    assert response.status_code == 422
    assert "detail" in response.json()


@pytest.mark.parametrize("body", [{"job_ids": ["1"]}, {"job_ids": [1.5]}, {"job_ids": [0]}, {}, {"job_ids": [2**63]}])
def test_body_failing_the_schema_is_a_default_422(client, body):
    response = client.post("/api/shipping-log", json=body)

    assert response.status_code == 422
    assert "detail" in response.json()


def test_ineligible_jobs_fail_the_whole_request_with_each_reason(client, session):
    eligible = _add_job(session, part_number="ELIG-OK")
    discarded = _add_job(session, part_number="ELIG-DISC", discarded_at=datetime(2026, 9, 1))
    superseded = _add_job(session, part_number="ELIG-SUP", superseded_at=datetime(2026, 9, 1))
    shipped = _add_job(session, part_number="ELIG-SHIP", status=JobStatus.shipped, shipped_at=date(2026, 9, 1))
    session.commit()
    not_found_id = shipped.id + 1000

    response = _generate(client, [eligible.id, discarded.id, superseded.id, shipped.id, not_found_id])

    assert response.status_code == 409
    assert response.headers["content-type"].startswith("application/json")
    assert response.json() == {
        "kind": "ineligible_jobs",
        "jobs": sorted([
            {"job_id": discarded.id, "reason": "discarded"},
            {"job_id": superseded.id, "reason": "superseded"},
            {"job_id": shipped.id, "reason": "shipped"},
            {"job_id": not_found_id, "reason": "not_found"},
        ], key=lambda job: job["job_id"]),
    }


def test_eligibility_precedence(session):
    both_discarded_and_superseded = _add_job(
        session, part_number="PREC-1", discarded_at=datetime(2026, 9, 1), superseded_at=datetime(2026, 9, 1),
    )
    both_superseded_and_shipped = _add_job(
        session, part_number="PREC-2", superseded_at=datetime(2026, 9, 1), status=JobStatus.shipped,
    )
    max_id = session.scalar(select(Job.id).order_by(Job.id.desc()).limit(1))

    eligibility = resolve_log_jobs(
        session, {both_discarded_and_superseded.id, both_superseded_and_shipped.id, max_id + 1},
    )

    assert eligibility.eligible == []
    assert {job.job_id: job.reason for job in eligibility.ineligible} == {
        both_discarded_and_superseded.id: IneligibleReason.DISCARDED,
        both_superseded_and_shipped.id: IneligibleReason.SUPERSEDED,
        max_id + 1: IneligibleReason.NOT_FOUND,
    }


def test_eligible_and_ineligible_partition_the_request(session):
    planned = _add_job(session, part_number="PART-1")
    shipped = _add_job(session, part_number="PART-2", status=JobStatus.shipped)

    eligibility = resolve_log_jobs(session, {planned.id, shipped.id})

    assert [job.id for job in eligibility.eligible] == [planned.id]
    assert [(job.job_id, job.reason) for job in eligibility.ineligible] == [(shipped.id, IneligibleReason.SHIPPED)]


def test_stamping_failure_is_a_500_without_job_data_in_the_log(
    client, session, monkeypatch, caplog, shipping_log_logger,
):
    from openpyxl.utils.exceptions import IllegalCharacterError
    from openpyxl.workbook.workbook import Workbook

    job = _add_job(session, part_number="SECRET-PART-777", notes="SECRET NOTE TEXT")
    session.commit()

    def _failing_save(self, target):
        raise IllegalCharacterError("SECRET NOTE TEXT cannot be used in worksheets.")

    monkeypatch.setattr(Workbook, "save", _failing_save)
    caplog.set_level(logging.ERROR, logger="backend.app.api.shipping_log")

    response = _generate(client, [job.id])

    assert response.status_code == 500
    assert response.json() == {"kind": "internal"}
    failure_lines = [r.getMessage() for r in caplog.records if r.name == "backend.app.api.shipping_log"]
    assert len(failure_lines) == 1
    assert "jobs=1" in failure_lines[0]
    assert _bundled_template().sha256 in failure_lines[0]
    assert "IllegalCharacterError" in failure_lines[0]
    assert "SECRET" not in caplog.text
    assert all(r.exc_info is None for r in caplog.records)


def test_stamp_workbook_wraps_openpyxl_failures(session, monkeypatch):
    from openpyxl.workbook.workbook import Workbook

    template = _bundled_template()
    boxes = [render_box(_add_job(session, part_number="WRAP-1"), Settings())]
    plan = plan_shipping_log_layout(template, boxes, Settings())

    def _failing_save(self, target):
        raise OSError("disk full")

    monkeypatch.setattr(Workbook, "save", _failing_save)

    with pytest.raises(WorkbookStampingError) as raised:
        stamp_workbook(template, boxes, plan)

    assert (raised.value.stage, raised.value.cause_type) == ("save", "OSError")
    assert raised.value.__suppress_context__ is True


def test_database_failure_is_a_500_with_counts_only(client, session, monkeypatch, caplog, shipping_log_logger):
    from sqlalchemy.exc import OperationalError

    def _failing_resolve(session, requested_ids):
        raise OperationalError("SELECT ...", {}, Exception("database is locked"))

    monkeypatch.setattr(shipping_log_service, "resolve_log_jobs", _failing_resolve)
    caplog.set_level(logging.ERROR, logger="backend.app.api.shipping_log")

    response = _generate(client, [5, 6, 6])

    assert response.status_code == 500
    assert response.json() == {"kind": "internal"}
    failure_lines = [r.getMessage() for r in caplog.records if r.name == "backend.app.api.shipping_log"]
    assert len(failure_lines) == 1
    assert "jobs=2" in failure_lines[0]
    assert "database is locked" in failure_lines[0]


# ---- §6.3 generation: response and side effects -------------------------------


def test_filename_uses_display_timezone_and_length_matches_body(client, session):
    job = _add_job(session, part_number="NAME-1")
    session.commit()
    _use_settings(client, display_timezone="Asia/Tokyo")
    client.app.dependency_overrides[get_wall_clock] = lambda: (lambda zone: PINNED_UTC.astimezone(zone))

    response = _generate(client, [job.id])

    assert response.headers["content-disposition"] == 'attachment; filename="Shipping_Log_20260916_213045.xlsx"'
    assert response.headers["content-type"] == shipping_log_service.SHIPPING_LOG_MEDIA_TYPE
    assert int(response.headers["content-length"]) == len(response.content)
    assert response.headers.get("content-encoding") in (None, "identity")


def test_dev_cors_exposes_the_filename_and_clipped_headers(client, session):
    job = _add_job(session, part_number="CORS-1", notes="x" * 40_000)
    session.commit()

    response = client.post(
        "/api/shipping-log", json={"job_ids": [job.id]}, headers={"Origin": "http://localhost:5173"},
    )

    exposed = {name.strip().lower() for name in response.headers["access-control-expose-headers"].split(",")}
    assert {"content-disposition", "x-shipping-log-clipped"} <= exposed


def test_default_display_timezone_is_pacific(client, session):
    job = _add_job(session, part_number="NAME-2")
    session.commit()
    client.app.dependency_overrides[get_wall_clock] = lambda: (lambda zone: PINNED_UTC.astimezone(zone))

    response = _generate(client, [job.id])

    assert 'filename="Shipping_Log_20260916_053045.xlsx"' in response.headers["content-disposition"]


def test_generation_writes_nothing_to_the_database(client, session, engine):
    jobs = [_add_job(session, part_number=f"RO-{i}", notes="note") for i in range(3)]
    session.commit()

    def _snapshot() -> dict[str, list]:
        with engine.connect() as connection:
            return {
                table.name: [tuple(row) for row in connection.execute(text(f'SELECT * FROM "{table.name}"'))]
                for table in Base.metadata.sorted_tables
            }

    before = _snapshot()
    response = _generate(client, [job.id for job in jobs])
    after = _snapshot()

    assert response.status_code == 200
    assert after == before


def test_max_selection_generates_within_the_client_timeout_budget(client, session):
    max_jobs = Settings().shipping_log_max_jobs
    notes = ("A" * 80 + "\n") * 8 + "A" * 7
    assert len(notes) == 655
    jobs = [_add_job(session, part_number=f"PERF-{i:03d}", notes=notes) for i in range(max_jobs)]
    session.commit()

    started = time.perf_counter()
    response = _generate(client, [job.id for job in jobs])
    elapsed = time.perf_counter() - started

    print(f"\nshipping log: {max_jobs} boxes with {len(notes)}-character notes in {elapsed:.2f}s")
    assert response.status_code == 200
    assert elapsed < 25.0
    sheet = _sheet_of(response.content)
    assert len(sheet.merged_cells.ranges) == max_jobs
    assert sheet.cell(row=_box_start(max_jobs - 1) + 1, column=1).value == f"PERF-{max_jobs - 1:03d}"


# ---- §3.6 CSV export filename ------------------------------------------------


def test_history_export_filename_uses_display_timezone(client):
    _use_settings(client, display_timezone="Asia/Tokyo")
    client.app.dependency_overrides[get_wall_clock] = lambda: (lambda zone: PINNED_UTC.astimezone(zone))

    response = client.get("/api/jobs/history/export.csv", params=[("column", "job"), ("delimiter", "comma")])

    assert response.status_code == 200
    assert response.headers["content-disposition"] == 'attachment; filename="history-export-20260916-213045.csv"'


def test_invalid_display_timezone_fails_settings_construction():
    from pydantic import ValidationError

    for zone in ("Mars/Olympus_Mons", "America", "", "../etc/passwd"):
        with pytest.raises(ValidationError):
            Settings(display_timezone=zone)


def test_notes_row_ceiling_cannot_be_raised_above_excel_limit():
    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        Settings(shipping_log_notes_row_max_points=410)
    assert Settings(shipping_log_notes_row_max_points=300).shipping_log_notes_row_max_points == 300
