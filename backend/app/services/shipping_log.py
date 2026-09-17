"""Shipping log generation (Phase 33).

A shipping log is one copy of the bundled template's six-row box per selected
planned job. The generator's only coupling to the template is
ShippingLogTemplateLayout: it never searches the sheet for cells.

Only the template's bytes are cached (on app.state). Every request loads its own
Workbook from them, because a shared Workbook would be mutable state across the
threadpool. Generating a log is read-only: no database row is written.
"""
from __future__ import annotations

import enum
import hashlib
import io
import logging
import math
from collections.abc import Callable
from copy import copy
from dataclasses import dataclass
from datetime import datetime, tzinfo
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, MergedCell
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload, selectinload

from ..config import Settings, bundled_resource_path
from ..models import Job, JobStatus
from ..schemas import ShippingLogCandidate, ShippingLogCandidateList
from .history_export import neutralise_formula_prefix, strip_operator_markup
from .jobs import build_shipping_query, count_shipping_population

logger = logging.getLogger(__name__)

SHIPPING_LOG_TEMPLATE_FILENAME = "shipping_log_template.xlsx"
SHIPPING_LOG_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Excel's per-cell limit, counted in UTF-16 code units. Past it the file won't open.
EXCEL_CELL_MAX_CHARS = 32767
TRUNCATION_MARKER = "… [truncated]"
LINE_FEED = "\n"
# openpyxl's SheetFormatProperties default, for a template that omits defaultRowHeight.
FALLBACK_DEFAULT_ROW_HEIGHT_POINTS = 15.0

WallClock = Callable[[tzinfo], datetime]


# ---- template layout ---------------------------------------------------------


@dataclass(frozen=True)
class ShippingLogTemplateLayout:
    """Where things are in the template. Fixed in code, not configuration.

    It describes a file that ships with the code; a configurable layout could
    disagree with the bundled template with no way to detect it.
    """

    sheet_title: str
    box_rows: int
    box_columns: int
    header_expectations: tuple[tuple[str, str], ...]
    notes_merge: str
    part_number_cell: str
    quantity_cell: str

    @property
    def notes_cell(self) -> str:
        notes_range = CellRange(self.notes_merge)
        return f"{get_column_letter(notes_range.min_col)}{notes_range.min_row}"

    @property
    def notes_row(self) -> int:
        return CellRange(self.notes_merge).min_row

    @property
    def notes_spanned_rows(self) -> range:
        notes_range = CellRange(self.notes_merge)
        return range(notes_range.min_row, notes_range.max_row + 1)


SHIPPING_LOG_TEMPLATE_LAYOUT = ShippingLogTemplateLayout(
    sheet_title="SHIPPING TEMPLATE",
    box_rows=6,
    box_columns=5,
    header_expectations=(
        ("A1", "B #"), ("B1", "QTY"), ("C1", "BAL"), ("D1", "NOTES"),
        ("A3", "SHIP TYPE"), ("A4", "INTL"), ("A5", "UPS"), ("A6", "FEDEX"),
    ),
    notes_merge="E2:E5",
    part_number_cell="A2",
    quantity_cell="B2",
)


class TemplateUnavailableKind(str, enum.Enum):
    MISSING = "missing"
    UNREADABLE = "unreadable"
    NOT_A_WORKBOOK = "not_a_workbook"
    SHEET_MISSING = "sheet_missing"
    HEADER_MISMATCH = "header_mismatch"
    NOTES_MERGE_MISSING = "notes_merge_missing"
    EXTRA_CONTENT_BELOW_BOX = "extra_content_below_box"


@dataclass(frozen=True)
class TemplateReady:
    """Validated template bytes plus the measurements stamping needs.

    Heights are in points; rows with no explicit height count at the sheet
    default row height.
    """

    content: bytes
    sha256: str
    layout: ShippingLogTemplateLayout
    box_height_points: float
    notes_row_height_points: float
    notes_available_points: float
    default_row_height_points: float


@dataclass(frozen=True)
class TemplateUnavailable:
    """Why the template can't be used. Logged at startup; never sent to a client."""

    kind: TemplateUnavailableKind
    sha256: str | None = None
    cell: str | None = None
    found: str | None = None

    def describe(self) -> str:
        if self.kind is TemplateUnavailableKind.HEADER_MISMATCH:
            return f"{self.kind.value} at {self.cell} (found {self.found!r})"
        if self.kind is TemplateUnavailableKind.EXTRA_CONTENT_BELOW_BOX:
            return f"{self.kind.value} at {self.cell}"
        return self.kind.value


def load_shipping_log_template(
    path: Path,
    layout: ShippingLogTemplateLayout,
) -> TemplateReady | TemplateUnavailable:
    """Load and structurally validate the template once, before requests are served.

    Post:   TemplateReady holding the raw bytes, their SHA-256 and the measured
            heights; or TemplateUnavailable when the file is missing, unreadable,
            not a workbook, or any layout expectation fails.
            EXTRA_CONTENT_BELOW_BOX rejects a template edited to hold more than
            one box: stamping assumes rows below the box are empty.
    Raises: never.
    """
    try:
        content = path.read_bytes()
    except FileNotFoundError:
        return TemplateUnavailable(TemplateUnavailableKind.MISSING)
    except OSError:
        return TemplateUnavailable(TemplateUnavailableKind.UNREADABLE)
    sha256 = hashlib.sha256(content).hexdigest()

    try:
        workbook = load_workbook(io.BytesIO(content))
    except Exception:
        return TemplateUnavailable(TemplateUnavailableKind.NOT_A_WORKBOOK, sha256)

    if layout.sheet_title not in workbook.sheetnames:
        return TemplateUnavailable(TemplateUnavailableKind.SHEET_MISSING, sha256)
    sheet = workbook[layout.sheet_title]

    for cell_ref, expected_label in layout.header_expectations:
        found = sheet[cell_ref].value
        if found != expected_label:
            return TemplateUnavailable(
                TemplateUnavailableKind.HEADER_MISMATCH, sha256,
                cell=cell_ref, found=None if found is None else str(found),
            )

    if layout.notes_merge not in {merged.coord for merged in sheet.merged_cells.ranges}:
        return TemplateUnavailable(TemplateUnavailableKind.NOTES_MERGE_MISSING, sha256)

    content_below = _first_content_below_box(sheet, layout.box_rows)
    if content_below is not None:
        return TemplateUnavailable(
            TemplateUnavailableKind.EXTRA_CONTENT_BELOW_BOX, sha256, cell=content_below,
        )

    return TemplateReady(
        content=content,
        sha256=sha256,
        layout=layout,
        box_height_points=sum(_row_height_points(sheet, row) for row in range(1, layout.box_rows + 1)),
        notes_row_height_points=_row_height_points(sheet, layout.notes_row),
        notes_available_points=sum(_row_height_points(sheet, row) for row in layout.notes_spanned_rows),
        default_row_height_points=_default_row_height_points(sheet),
    )


def load_bundled_shipping_log_template() -> TemplateReady | TemplateUnavailable:
    """Load the template bundled with this build, logging the outcome once.

    An unavailable template does not stop startup: the rest of the app is
    unaffected, generate answers 503 and the candidates response says so.
    """
    path = bundled_resource_path(SHIPPING_LOG_TEMPLATE_FILENAME)
    template = load_shipping_log_template(path, SHIPPING_LOG_TEMPLATE_LAYOUT)
    if isinstance(template, TemplateUnavailable):
        logger.error(
            "Shipping log template unavailable: %s (path=%s, sha256=%s). "
            "Generating shipping logs will answer 503 until the build is fixed.",
            template.describe(), path, template.sha256 or "unreadable",
        )
    else:
        logger.info("Shipping log template loaded (sha256=%s)", template.sha256)
    return template


def _default_row_height_points(sheet: Worksheet) -> float:
    return sheet.sheet_format.defaultRowHeight or FALLBACK_DEFAULT_ROW_HEIGHT_POINTS


def _row_height_points(sheet: Worksheet, row: int) -> float:
    # .get, not [row]: row_dimensions creates an entry on subscript access.
    dimension = sheet.row_dimensions.get(row)
    if dimension is None or dimension.height is None:
        return _default_row_height_points(sheet)
    return dimension.height


def _first_content_below_box(sheet: Worksheet, box_rows: int) -> str | None:
    for merged in sheet.merged_cells.ranges:
        if merged.max_row > box_rows:
            return merged.coord
    for row in sheet.iter_rows(min_row=box_rows + 1):
        for cell in row:
            if cell.value is not None:
                return cell.coordinate
    return None


# ---- candidates --------------------------------------------------------------


def list_shipping_log_candidates(
    session: Session,
    settings: Settings,
    template: TemplateReady | TemplateUnavailable,
) -> ShippingLogCandidateList:
    """List the jobs a shipping log may be generated from.

    Post:   at most settings.shipping_log_candidate_max candidates, in
            build_shipping_query order; total is the unbounded population count
            and truncated == (total > len(candidates)). Both reads share the
            session's transaction, so they see one snapshot.
    Raises: SQLAlchemyError (propagated).
    """
    total = count_shipping_population(session)
    jobs = session.scalars(
        build_shipping_query()
        .options(selectinload(Job.assembly), selectinload(Job.customer))
        .limit(settings.shipping_log_candidate_max)
    ).all()
    candidates = [
        ShippingLogCandidate(
            job_id=job.id,
            part_number=job.assembly.part_number,
            split_suffix=job.split_suffix,
            build_type=job.build_type,
            repeat_reference=job.repeat_reference,
            build_qualifier=job.build_qualifier,
            quantity=job.quantity,
            resolved_ship_date=job.resolved_ship_date,
            ship_date_text=job.ship_date_text,
            customer_name=job.customer.name,
        )
        for job in jobs
    ]
    return ShippingLogCandidateList(
        candidates=candidates,
        total=total,
        truncated=total > len(candidates),
        max_jobs_per_log=settings.shipping_log_max_jobs,
        template_ready=isinstance(template, TemplateReady),
    )


# ---- eligibility -------------------------------------------------------------


class IneligibleReason(str, enum.Enum):
    NOT_FOUND = "not_found"
    DISCARDED = "discarded"
    SUPERSEDED = "superseded"
    SHIPPED = "shipped"


@dataclass(frozen=True)
class IneligibleJob:
    job_id: int
    reason: IneligibleReason


@dataclass(frozen=True)
class JobExclusionRow:
    status: JobStatus
    discarded_at: datetime | None
    superseded_at: datetime | None


@dataclass(frozen=True)
class JobEligibility:
    eligible: list[Job]
    ineligible: list[IneligibleJob]


def resolve_log_jobs(session: Session, requested_ids: set[int]) -> JobEligibility:
    """Classify every requested ID against the Shipping-view population.

    Pre:    requested_ids is non-empty and no larger than shipping_log_max_jobs,
            which keeps the IN lists below SQLite's bound-parameter limit.
    Post:   eligible holds every requested job present in build_shipping_query(),
            in that query's order, with assembly loaded; ineligible holds one
            entry per remaining ID, ordered by ID. The two partition requested_ids.
    Raises: SQLAlchemyError (propagated).
    """
    eligible = list(session.scalars(
        build_shipping_query()
        .where(Job.id.in_(requested_ids))
        .options(joinedload(Job.assembly))
    ).all())

    missed_ids = requested_ids - {job.id for job in eligible}
    exclusion_rows: dict[int, JobExclusionRow] = {}
    if missed_ids:
        for job_id, status, discarded_at, superseded_at in session.execute(
            select(Job.id, Job.status, Job.discarded_at, Job.superseded_at)
            .where(Job.id.in_(missed_ids))
        ):
            exclusion_rows[job_id] = JobExclusionRow(status, discarded_at, superseded_at)

    ineligible = [
        IneligibleJob(job_id, classify_ineligible(exclusion_rows.get(job_id)))
        for job_id in sorted(missed_ids)
    ]
    return JobEligibility(eligible=eligible, ineligible=ineligible)


def classify_ineligible(exclusion_row: JobExclusionRow | None) -> IneligibleReason:
    """The reason a requested ID is outside the Shipping-view population.

    First match wins: NOT_FOUND, DISCARDED, SUPERSEDED, SHIPPED. Discarded
    outranks the others because it is the state the operator most likely caused
    and can undo from the Discarded drawer.
    status == shipped is the only exclusion left in build_shipping_query once
    the first three are ruled out, so the last branch needs no test of its own.
    Raises: never.
    """
    if exclusion_row is None:
        return IneligibleReason.NOT_FOUND
    if exclusion_row.discarded_at is not None:
        return IneligibleReason.DISCARDED
    if exclusion_row.superseded_at is not None:
        return IneligibleReason.SUPERSEDED
    return IneligibleReason.SHIPPED


# ---- box rendering -----------------------------------------------------------


@dataclass(frozen=True)
class TextCellValue:
    """Text that is safe to put in a cell. Construct only via text_cell_value."""

    text: str


def text_cell_value(raw_text: str) -> TextCellValue:
    """The only constructor of TextCellValue.

    Pre:    raw_text is under EXCEL_CELL_MAX_CHARS UTF-16 units (render_box
            truncates first, leaving room for the one-character prefix).
    Post:   characters XML 1.0 can't carry are removed (openpyxl would otherwise
            refuse the whole workbook), then neutralise_formula_prefix is applied.
    Raises: never.
    """
    return TextCellValue(neutralise_formula_prefix(ILLEGAL_CHARACTERS_RE.sub("", raw_text)))


def write_text_cell(cell, value: TextCellValue) -> None:
    """The only path by which a TextCellValue reaches a worksheet.

    Post:   empty text leaves the cell blank. Otherwise cell.value == value.text
            and cell.data_type == 's', assigned after the value, so openpyxl's
            formula inference on '=' (and error inference on '#N/A') is
            overwritten and no <f> element is ever written.
    Raises: never.
    """
    if not value.text:
        cell.value = None
        return
    cell.value = value.text
    cell.data_type = "s"


@dataclass(frozen=True)
class ShippingLogBox:
    job_id: int
    part_number: TextCellValue
    quantity: int
    notes: TextCellValue
    notes_line_estimate: int
    notes_clipped: bool


def render_box(job: Job, settings: Settings) -> ShippingLogBox:
    """The values written into one box, all cell-safe.

    Pre:    job.assembly is loaded.
    Post:   notes are assembly.base_mfg_notes (Decision 8) with operator markup
            stripped and lines joined by LINE_FEED; notes over the cell limit are
            cut with TRUNCATION_MARKER and notes_clipped is set.
            notes_line_estimate = sum over lines of
            ceil(max(1, len(line)) / shipping_log_notes_chars_per_line).
    Raises: never.
    """
    notes_text = LINE_FEED.join(strip_operator_markup(job.assembly.base_mfg_notes))
    notes_clipped = False
    notes_budget = EXCEL_CELL_MAX_CHARS - 1
    if _utf16_units(notes_text) > notes_budget:
        kept_units = notes_budget - _utf16_units(TRUNCATION_MARKER)
        notes_text = _truncate_to_utf16_units(notes_text, kept_units) + TRUNCATION_MARKER
        notes_clipped = True

    chars_per_line = settings.shipping_log_notes_chars_per_line
    notes_line_estimate = sum(
        math.ceil(max(1, len(line)) / chars_per_line)
        for line in notes_text.split(LINE_FEED)
    ) if notes_text else 0

    return ShippingLogBox(
        job_id=job.id,
        part_number=text_cell_value(job.assembly.part_number),
        quantity=job.quantity,
        notes=text_cell_value(notes_text),
        notes_line_estimate=notes_line_estimate,
        notes_clipped=notes_clipped,
    )


def _utf16_units(text: str) -> int:
    return len(text.encode("utf-16-le", errors="surrogatepass")) // 2


def _truncate_to_utf16_units(text: str, units: int) -> str:
    # A surrogate pair cut in half decodes to nothing rather than to a lone surrogate.
    return text.encode("utf-16-le", errors="surrogatepass")[: 2 * units].decode("utf-16-le", errors="ignore")


# ---- layout planning ---------------------------------------------------------


@dataclass(frozen=True)
class BoxPlacement:
    start_row: int
    notes_row_points: float
    height_points: float
    notes_height_clipped: bool


@dataclass(frozen=True)
class ShippingLogLayoutPlan:
    placements: list[BoxPlacement]
    row_breaks_after: list[int]


def plan_shipping_log_layout(
    template: TemplateReady,
    boxes: list[ShippingLogBox],
    settings: Settings,
) -> ShippingLogLayoutPlan:
    """Where each box goes, how tall its notes row is, and where pages break.

    Pre:    boxes is non-empty.
    Post:   box k starts at row 1 + k * (box_rows + shipping_log_box_gap_rows).
            A merged cell is never auto-fit by Excel, so the notes row grows by
            however far the estimated notes height exceeds the merged rows'
            height, capped at shipping_log_notes_row_max_points; hitting the cap
            sets notes_height_clipped.
            Pages are packed greedily on each box's printed height, gap rows
            included at the sheet default row height. A break goes after the
            last row of the previous box, so the gap rows start the new page and
            are counted there (the cost: one blank gap at the top of every later
            page). No box is ever split; a box taller than a page gets its own.
    Raises: never.
    """
    layout = template.layout
    gap_rows = settings.shipping_log_box_gap_rows
    rows_per_box = layout.box_rows + gap_rows

    placements: list[BoxPlacement] = []
    for index, box in enumerate(boxes):
        required_points = box.notes_line_estimate * settings.shipping_log_notes_line_points
        growth_points = max(0.0, required_points - template.notes_available_points)
        wanted_row_points = template.notes_row_height_points + growth_points
        notes_row_points = min(wanted_row_points, settings.shipping_log_notes_row_max_points)
        placements.append(BoxPlacement(
            start_row=1 + index * rows_per_box,
            notes_row_points=notes_row_points,
            height_points=template.box_height_points - template.notes_row_height_points + notes_row_points,
            notes_height_clipped=wanted_row_points > settings.shipping_log_notes_row_max_points,
        ))

    gap_points = gap_rows * template.default_row_height_points
    page_body_points = settings.shipping_log_page_body_points
    row_breaks_after: list[int] = []
    running_points = placements[0].height_points
    for previous, placement in zip(placements, placements[1:]):
        if running_points + gap_points + placement.height_points <= page_body_points:
            running_points += gap_points + placement.height_points
        else:
            row_breaks_after.append(previous.start_row + layout.box_rows - 1)
            running_points = gap_points + placement.height_points

    return ShippingLogLayoutPlan(placements=placements, row_breaks_after=row_breaks_after)


# ---- stamping ----------------------------------------------------------------


class WorkbookStampingError(Exception):
    """openpyxl failed to load, stamp or save the validated template bytes.

    Carries the failing stage and the cause's class name only. openpyxl messages
    can quote cell text (IllegalCharacterError does) and note text must not
    reach the log, so the cause is not chained.
    """

    def __init__(self, stage: str, cause_type: str) -> None:
        super().__init__(f"shipping log stamping failed during {stage}: {cause_type}")
        self.stage = stage
        self.cause_type = cause_type


def stamp_workbook(
    template: TemplateReady,
    boxes: list[ShippingLogBox],
    plan: ShippingLogLayoutPlan,
) -> bytes:
    """Stamp one copy of the template box per ShippingLogBox and serialise.

    Pre:    boxes is non-empty; plan came from plan_shipping_log_layout for
            these boxes.
    Post:   two passes over a fresh workbook. The copy pass copies rows
            1..box_rows (values, styles, row heights, merges) to every later
            box while those rows still hold only template content. The value
            pass then writes each box's part number, quantity and notes and sets
            its notes row height. Writing any value before the copy pass ends
            would duplicate job 0's values into every later box, and a later
            box with empty notes would print job 0's notes.
            Column widths, orientation and margins are the template's.
    Raises: WorkbookStampingError.
    """
    layout = template.layout
    stage = "load"
    try:
        workbook = load_workbook(io.BytesIO(template.content))
        sheet = workbook[layout.sheet_title]

        stage = "copy"
        for placement in plan.placements[1:]:
            _copy_template_box(sheet, layout, row_offset=placement.start_row - 1)

        stage = "values"
        part_number_row, part_number_column = coordinate_to_tuple(layout.part_number_cell)
        quantity_row, quantity_column = coordinate_to_tuple(layout.quantity_cell)
        notes_row, notes_column = coordinate_to_tuple(layout.notes_cell)
        for box, placement in zip(boxes, plan.placements):
            row_offset = placement.start_row - 1
            write_text_cell(
                sheet.cell(row=part_number_row + row_offset, column=part_number_column),
                box.part_number,
            )
            sheet.cell(row=quantity_row + row_offset, column=quantity_column).value = box.quantity
            write_text_cell(sheet.cell(row=notes_row + row_offset, column=notes_column), box.notes)
            sheet.row_dimensions[notes_row + row_offset].height = placement.notes_row_points

        stage = "breaks"
        for last_row in plan.row_breaks_after:
            sheet.row_breaks.append(Break(id=last_row))

        stage = "save"
        buffer = io.BytesIO()
        workbook.save(buffer)
    except Exception as exc:
        raise WorkbookStampingError(stage, type(exc).__name__) from None
    return buffer.getvalue()


def _copy_template_box(sheet: Worksheet, layout: ShippingLogTemplateLayout, row_offset: int) -> None:
    # Merge first: merge_cells replaces the covered cells with fresh MergedCells,
    # so styles copied before it would be discarded.
    for merged in list(sheet.merged_cells.ranges):
        if merged.max_row <= layout.box_rows:
            sheet.merge_cells(
                start_row=merged.min_row + row_offset, start_column=merged.min_col,
                end_row=merged.max_row + row_offset, end_column=merged.max_col,
            )

    for source_row in range(1, layout.box_rows + 1):
        target_row = source_row + row_offset
        source_dimension = sheet.row_dimensions.get(source_row)
        if source_dimension is not None:
            target_dimension = sheet.row_dimensions[target_row]
            target_dimension.height = source_dimension.height
            target_dimension.thickBot = source_dimension.thickBot
            target_dimension.thickTop = source_dimension.thickTop
        for column in range(1, layout.box_columns + 1):
            source = sheet.cell(row=source_row, column=column)
            target = sheet.cell(row=target_row, column=column)
            if not isinstance(target, MergedCell):
                target.value = source.value
            target._style = copy(source._style)


# ---- generation --------------------------------------------------------------


@dataclass(frozen=True)
class TemplateUnavailableFailure:
    kind: TemplateUnavailableKind


@dataclass(frozen=True)
class SelectionSizeInvalid:
    requested: int
    max: int


@dataclass(frozen=True)
class IneligibleJobs:
    jobs: list[IneligibleJob]


# Outcomes the operator can act on. Database and stamping faults are not in this
# union: they propagate as exceptions to the 500 mapping, because the client can
# do nothing with them except retry.
ShippingLogFailure = TemplateUnavailableFailure | SelectionSizeInvalid | IneligibleJobs


@dataclass(frozen=True)
class ShippingLogFile:
    content: bytes
    filename: str
    clipped_job_ids: list[int]


def shipping_log_filename(local_now: datetime) -> str:
    return f"Shipping_Log_{local_now:%Y%m%d_%H%M%S}.xlsx"


def generate_shipping_log(
    job_ids: list[int],
    session: Session,
    settings: Settings,
    template: TemplateReady | TemplateUnavailable,
    now_in: WallClock,
) -> ShippingLogFile | ShippingLogFailure:
    """Generate a shipping log workbook for the selected planned jobs.

    Post:   on success, the workbook bytes, the server-chosen filename in
            display_timezone, and the IDs whose notes won't print in full (cut
            at the cell limit or at the row-height ceiling). Every domain
            failure is returned as a value. The request is all-or-nothing: any
            ineligible ID fails it. No database row is written. Box order is
            the Shipping-view order, not the order IDs were submitted.
    Raises: SQLAlchemyError (propagated from resolve_log_jobs);
            WorkbookStampingError (propagated from stamp_workbook).
    """
    if isinstance(template, TemplateUnavailable):
        return TemplateUnavailableFailure(template.kind)

    requested_ids = set(job_ids)
    if not requested_ids or len(requested_ids) > settings.shipping_log_max_jobs:
        return SelectionSizeInvalid(requested=len(requested_ids), max=settings.shipping_log_max_jobs)

    eligibility = resolve_log_jobs(session, requested_ids)
    if eligibility.ineligible:
        return IneligibleJobs(eligibility.ineligible)

    boxes = [render_box(job, settings) for job in eligibility.eligible]
    plan = plan_shipping_log_layout(template, boxes, settings)
    content = stamp_workbook(template, boxes, plan)
    return ShippingLogFile(
        content=content,
        filename=shipping_log_filename(now_in(ZoneInfo(settings.display_timezone))),
        clipped_job_ids=[
            box.job_id
            for box, placement in zip(boxes, plan.placements)
            if box.notes_clipped or placement.notes_height_clipped
        ],
    )
